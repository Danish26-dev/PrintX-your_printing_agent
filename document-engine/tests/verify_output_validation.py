"""Verify output validation: passes on a correct edit, fails on a corrupt file,
and does not over-assert when text was not found (0 replacements).

Run:  venv/Scripts/python.exe -m tests.verify_output_validation
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

from app.editing.schema import OperationPlan
from app.service import process_document
from app.storage.files import create_job
from app.validation import validate_output, OutputValidationError


def main() -> int:
    failures = []

    # 1. Good edit -> validation ok, find removed / replace present.
    job = create_job()
    sample = job.path("sample.docx")
    d = Document()
    d.add_paragraph("Hello Rahul, regards Rahul.")
    d.save(str(sample))

    plan = OperationPlan.model_validate(
        {"operations": [{"type": "replace_text", "find": "Rahul", "replace": "Danish"}]}
    )
    outcome = process_document(job=job, source_path=sample, filename="sample.docx", plan=plan)
    v = outcome.validation
    if not v["ok"]:
        failures.append("valid edit should pass validation")
    if v["checks"].get("find_removed") is not True:
        failures.append("find_removed should be True")
    if v["checks"].get("replace_present") is not True:
        failures.append("replace_present should be True")

    # 2. Corrupt/empty output -> hard failure.
    job2 = create_job()
    empty = job2.path("empty.docx")
    empty.write_bytes(b"")  # not a valid docx
    try:
        validate_output(empty, "docx", plan, executed_replacements=1)
        failures.append("empty output should raise OutputValidationError")
    except OutputValidationError:
        pass

    # 3. Text not found (0 replacements) -> no over-assertion, still ok.
    job3 = create_job()
    s3 = job3.path("s3.docx")
    d3 = Document()
    d3.add_paragraph("No target here.")
    d3.save(str(s3))
    plan3 = OperationPlan.model_validate(
        {"operations": [{"type": "replace_text", "find": "Nonexistent", "replace": "X"}]}
    )
    outcome3 = process_document(job=job3, source_path=s3, filename="s3.docx", plan=plan3)
    if not outcome3.validation["ok"]:
        failures.append("0-replacement edit should still validate ok")
    if outcome3.change_summary["total_replacements"] != 0:
        failures.append("expected 0 replacements")

    if failures:
        print("FAILED:")
        for f in failures:
            print("  -", f)
        return 1
    print("ALL OUTPUT-VALIDATION CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
