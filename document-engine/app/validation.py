"""Output validation for processed documents.

After the deterministic edit + render, we verify the output is sound before the
engine reports success. This is a safety net so a corrupted/empty output is
never presented as a valid result. It is honest: it reports what it could and
could not verify, and raises for clear failures.

Checks performed:
  - output file exists and is non-empty
  - output opens successfully for its format (docx / pdf)
  - the edited text no longer contains any replaced `find` strings, and does
    contain the `replace` strings (when a replacement actually matched)
  - page count is reported (and, for PDF round-trips, we note it cannot be
    guaranteed identical — we never claim pixel fidelity)
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from .editing.schema import OperationPlan, ReplaceTextOperation


class OutputValidationError(Exception):
    """Raised when the produced output fails a hard validation check."""

    code = "output_validation_failed"


@dataclass
class ValidationReport:
    ok: bool
    output_format: str
    page_count: int | None = None
    checks: dict = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)
    # Which edit strategy produced the output: native_pdf | conversion | docx_direct
    strategy: str | None = None
    # Structural comparison against the original (PDF only). This is REPORTED,
    # not enforced: unexpected differences are flagged for review, because some
    # legitimate edits legitimately reflow the document.
    structure: dict = field(default_factory=dict)
    # True when a difference was detected that a human should review (e.g. the
    # page count changed, or blank pages appeared). Does NOT by itself fail the
    # job — it is surfaced so the UI/caller can warn the user honestly.
    review_recommended: bool = False

    def as_dict(self) -> dict:
        return {
            "ok": self.ok,
            "output_format": self.output_format,
            "page_count": self.page_count,
            "strategy": self.strategy,
            "checks": self.checks,
            "structure": self.structure,
            "review_recommended": self.review_recommended,
            "notes": self.notes,
        }


def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(t for t in parts if t)


def _pdf_text_and_pages(path: Path) -> tuple[str, int]:
    import fitz  # PyMuPDF

    text_parts: list[str] = []
    with fitz.open(str(path)) as doc:
        page_count = doc.page_count
        for page in doc:
            text_parts.append(page.get_text("text"))
    return "\n".join(text_parts), page_count


def _pdf_profile(path: Path) -> dict:
    """Collect structural facts about a PDF for before/after comparison."""
    import fitz  # PyMuPDF

    with fitz.open(str(path)) as doc:
        pages = doc.page_count
        dims = [(round(p.rect.width, 1), round(p.rect.height, 1)) for p in doc]
        images = sum(len(p.get_images(full=True)) for p in doc)
        blank = []
        for i, p in enumerate(doc):
            txt = p.get_text("text").strip()
            imgs = p.get_images(full=True)
            drawings = p.get_drawings()
            if len(txt) < 3 and len(imgs) == 0 and len(drawings) == 0:
                blank.append(i)
    return {"pages": pages, "dims": dims, "images": images, "blank_pages": blank}


# A page-count change beyond this ratio on a PDF round-trip is flagged for
# review (not forbidden). ~15% tolerance covers minor legitimate reflow.
_PAGE_COUNT_TOLERANCE = 0.15


def _compare_pdf_structure(original: Path, edited: Path) -> tuple[dict, list[str], bool]:
    """Compare the edited PDF against the original. Returns (structure, notes,
    review_recommended). Differences are REPORTED, never used to fail the job.
    """
    before = _pdf_profile(original)
    after = _pdf_profile(edited)
    notes: list[str] = []
    review = False

    page_delta = after["pages"] - before["pages"]
    if before["pages"] > 0:
        ratio = abs(page_delta) / before["pages"]
    else:
        ratio = 0.0

    # Dimensions: compare the most common page size before/after.
    def _dominant(dims: list[tuple]) -> tuple | None:
        if not dims:
            return None
        counts: dict[tuple, int] = {}
        for d in dims:
            counts[d] = counts.get(d, 0) + 1
        return max(counts, key=counts.get)

    dom_before = _dominant(before["dims"])
    dom_after = _dominant(after["dims"])
    dims_changed = False
    if dom_before and dom_after:
        dims_changed = (
            abs(dom_before[0] - dom_after[0]) > 1.5
            or abs(dom_before[1] - dom_after[1]) > 1.5
        )

    new_blanks = [b for b in after["blank_pages"] if b not in before["blank_pages"]]
    images_lost = after["images"] < before["images"] * 0.5 and before["images"] > 0

    if ratio > _PAGE_COUNT_TOLERANCE:
        review = True
        notes.append(
            f"Page count changed from {before['pages']} to {after['pages']} "
            f"({page_delta:+d}); layout may have reflowed. Flagged for review."
        )
    if dims_changed:
        review = True
        notes.append(
            f"Page dimensions changed from {dom_before} to {dom_after} points."
        )
    if new_blanks:
        review = True
        notes.append(f"Unexpected blank page(s) detected at index: {new_blanks}.")
    if images_lost:
        review = True
        notes.append(
            f"Image count dropped from {before['images']} to {after['images']}; "
            "some images may be missing."
        )

    structure = {
        "pages_before": before["pages"],
        "pages_after": after["pages"],
        "page_delta": page_delta,
        "dims_before": dom_before,
        "dims_after": dom_after,
        "dims_changed": dims_changed,
        "images_before": before["images"],
        "images_after": after["images"],
        "new_blank_pages": new_blanks,
        "layout_preserved": not review,
    }
    return structure, notes, review


def validate_output(
    output_path: Path,
    output_format: str,
    plan: OperationPlan,
    *,
    executed_replacements: int,
    original_pdf: Path | None = None,
    strategy: str | None = None,
) -> ValidationReport:
    """Validate the produced output. Raises OutputValidationError on hard
    failures; returns a report (ok=True) when checks pass.

    `original_pdf` (when the input was a PDF) enables a structural comparison
    (page count, dimensions, blank pages, image presence). Differences are
    FLAGGED for review via `review_recommended`; they do not fail the job,
    because some legitimate edits reflow the document.
    """
    report = ValidationReport(ok=True, output_format=output_format, strategy=strategy)

    # 1. Exists + non-empty.
    if not output_path.exists() or output_path.stat().st_size == 0:
        raise OutputValidationError("Output document is missing or empty.")
    report.checks["non_empty"] = True

    # 2. Opens for its format; extract text + page count.
    fmt = output_format.lower()
    try:
        if fmt == "docx":
            text = _docx_text(output_path)
            report.page_count = None  # DOCX has no fixed page count without rendering
            report.notes.append("DOCX page count is not fixed until rendered.")
        elif fmt == "pdf":
            text, pages = _pdf_text_and_pages(output_path)
            report.page_count = pages
        else:
            text = ""
            report.notes.append(f"No text validation for format '{fmt}'.")
    except Exception as exc:  # noqa: BLE001 - any open failure is a hard failure
        raise OutputValidationError(f"Output document could not be opened: {exc}") from exc
    report.checks["opens"] = True

    # 3. Replacement correctness (only when a replacement actually matched).
    #    If 0 replacements executed (text not found), we do NOT assert anything
    #    about find/replace — that is the caller's TEXT_NOT_FOUND concern.
    if executed_replacements > 0:
        leftover = []
        missing = []
        for op in plan.operations:
            if isinstance(op, ReplaceTextOperation):
                # The original `find` should be gone (unless `replace` contains it).
                if op.find in text and op.find not in op.replace:
                    leftover.append(op.find)
                # The `replace` text should now be present.
                if op.replace and op.replace not in text:
                    missing.append(op.replace)
        report.checks["find_removed"] = len(leftover) == 0
        report.checks["replace_present"] = len(missing) == 0
        if leftover:
            report.notes.append(f"Expected replaced text still present: {leftover}")
        if missing:
            report.notes.append(f"Expected replacement text not found: {missing}")
        # These are soft checks for PDF round-trips (conversion can alter text
        # segmentation); treat as hard only for DOCX where we control editing.
        if fmt == "docx" and (leftover or missing):
            report.ok = False
            raise OutputValidationError(
                "Replacement verification failed on the edited DOCX."
            )

    # 4. Structural / layout comparison (PDF only, when we have the original).
    #    This distinguishes TEXT-edit success from LAYOUT preservation. Any
    #    unexpected structural difference is flagged for review but does NOT
    #    fail the job — we report honestly rather than silently pass or fail.
    if fmt == "pdf" and original_pdf is not None and original_pdf.exists():
        try:
            structure, struct_notes, review = _compare_pdf_structure(original_pdf, output_path)
            report.structure = structure
            report.notes.extend(struct_notes)
            report.review_recommended = review
            report.checks["layout_preserved"] = structure.get("layout_preserved", False)
        except Exception as exc:  # noqa: BLE001 - comparison must never crash the job
            report.notes.append(f"Structural comparison unavailable: {exc}")

    return report
