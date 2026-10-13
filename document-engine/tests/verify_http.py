"""Live HTTP smoke test against a running Document Engine on 127.0.0.1:5001.

Uses only the stdlib (urllib) to avoid extra deps and shell-quoting issues.
Builds a multipart/form-data POST to /process, then downloads the result.

Run (with server running):  venv/Scripts/python.exe -m tests.verify_http
"""

from __future__ import annotations

import io
import json
import sys
import urllib.request
import uuid
from pathlib import Path

from docx import Document

BASE = "http://127.0.0.1:5001"


def make_docx(path: Path) -> None:
    d = Document()
    d.add_paragraph("Hello John from Acme Corp")
    d.save(str(path))


def multipart(fields: dict[str, str], files: dict[str, tuple[str, bytes]]) -> tuple[bytes, str]:
    boundary = "----printx" + uuid.uuid4().hex
    body = io.BytesIO()
    for name, value in fields.items():
        body.write(f"--{boundary}\r\n".encode())
        body.write(f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode())
        body.write(value.encode())
        body.write(b"\r\n")
    for name, (filename, content) in files.items():
        body.write(f"--{boundary}\r\n".encode())
        body.write(
            f'Content-Disposition: form-data; name="{name}"; filename="{filename}"\r\n'.encode()
        )
        body.write(b"Content-Type: application/octet-stream\r\n\r\n")
        body.write(content)
        body.write(b"\r\n")
    body.write(f"--{boundary}--\r\n".encode())
    return body.getvalue(), f"multipart/form-data; boundary={boundary}"


def post(url: str, body: bytes, content_type: str) -> dict:
    req = urllib.request.Request(url, data=body, method="POST")
    req.add_header("Content-Type", content_type)
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.loads(resp.read().decode())


def get_bytes(url: str) -> bytes:
    with urllib.request.urlopen(url, timeout=60) as resp:
        return resp.read()


def main() -> int:
    work = Path("_work")
    work.mkdir(exist_ok=True)
    sample = work / "http_test.docx"
    make_docx(sample)

    plan = json.dumps(
        {
            "operations": [
                {"type": "replace_text", "find": "John", "replace": "Jane"},
                {"type": "set_copies", "copies": 2},
            ]
        }
    )

    body, ctype = multipart(
        fields={"plan": plan},
        files={"file": ("http_test.docx", sample.read_bytes())},
    )
    result = post(f"{BASE}/process", body, ctype)
    print("--- /process response ---")
    print(json.dumps(result, indent=2))

    dl = get_bytes(f"{BASE}{result['download_url']}")
    out = work / "downloaded.docx"
    out.write_bytes(dl)
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    print("--- downloaded text ---")
    print(text)

    ok = (
        result["changes"]["copies"] == 2
        and result["changes"]["total_replacements"] == 1
        and "Jane" in text
        and "John" not in text
    )
    print("\n" + ("HTTP E2E PASSED" if ok else "HTTP E2E FAILED"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
