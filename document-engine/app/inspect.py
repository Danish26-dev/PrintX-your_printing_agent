"""Document inspection: extract plain text so the agent can reason about it.

The Strands agent needs to know what the document actually says to plan
sensible operations (e.g. which exact string to find/replace). This module
returns the document's text without modifying anything. PDFs are read directly
for text (fast) rather than round-tripping through DOCX.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from .conversion.engine import DocFormat, detect_format
from .conversion.pdf_to_docx import pdf_has_extractable_text
from .conversion.errors import ScannedPdfError
from .storage.files import Job


def _docx_text(docx_path: Path) -> str:
    doc = Document(str(docx_path))
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(t for t in parts if t is not None)


def _pdf_text(pdf_path: Path) -> str:
    import fitz  # PyMuPDF (bundled with pdf2docx)

    parts: list[str] = []
    with fitz.open(str(pdf_path)) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
    return "\n".join(parts)


def extract_text(job: Job, source_path: Path, filename: str) -> str:
    """Return the plain text of a supported document. Raises ScannedPdfError for
    image-based PDFs so the caller can respond clearly."""
    fmt = detect_format(filename)
    if fmt is DocFormat.DOCX:
        return _docx_text(source_path)
    if fmt is DocFormat.PDF:
        if not pdf_has_extractable_text(source_path):
            raise ScannedPdfError()
        return _pdf_text(source_path)
    return ""
